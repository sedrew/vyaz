#!/usr/bin/env python3
"""gen-textframe-fit.py — PowerPoint side of the fixed-rectangle round-trip.

Reads `textframe-fit.json` (written by `textframe-fit-run.ts`). Per frame it
puts TWO TextBoxes on one slide, same text / width / font:

  <name>       fixed at vyaz's width × content.height, wrap-only bodyPr, no
               autofit element — the box keeps exactly the size vyaz computed.
  <name>@fit   same width, height 1", "resize shape to fit text"
               (<a:spAutoFit/>) — PowerPoint rewrites its `cy` on save, so its
               height IS PowerPoint's own answer for the content height.

Compare the two: if they match, vyaz's `content.height` is what PowerPoint would
use. A gap means vyaz counts the last line's trailing half-leading that
PowerPoint trims (or vice-versa) — see RESULTS.md "last-line trailing leading".

  1. bun    scripts/office-metrics/textframe-fit-run.ts     # vyaz -> textframe-fit.json
  2. python3 scripts/office-metrics/gen-textframe-fit.py    # json -> textframe-fit.pptx
  3. open textframe-fit.pptx, SAVE (so PowerPoint recomputes every @fit cy),
     then File ▸ Export ▸ SVG per slide; send back the SVGs + the saved .pptx.

Needs python-pptx + Roboto installed. 1 vyaz unit == 1 pt.
"""
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR

HERE = __file__.rsplit("/", 1)[0]


def _textbox(slide, name, x, w_pt, h_pt, row, autofit):
    tb = slide.shapes.add_textbox(x, Inches(0.5), Pt(w_pt), Pt(h_pt))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT if autofit else None
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.line_spacing = row["line_spacing"]
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = row["text"]
    r.font.name = row["font_family"]
    r.font.size = Pt(row["font_size_pt"])


def add_slide(prs, row):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    w = row["width_pt"]
    _textbox(slide, row["name"], Inches(0.5), w, row["height_pt"], row, autofit=False)
    _textbox(slide, row["name"] + "@fit", Inches(0.5) + Pt(w) + Pt(36), w, 72, row, autofit=True)


def main():
    with open(f"{HERE}/textframe-fit.json") as fh:
        data = json.load(fh)
    rows = data["frames"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for row in rows:
        add_slide(prs, row)
    prs.save(f"{HERE}/textframe-fit.pptx")

    print(f"wrote textframe-fit.pptx ({len(rows)} slides, 2 boxes each) from textframe-fit.json")
    for row in rows:
        print(f"  {row['name']:20} fixed {row['width_pt']}×{row['height_pt']}pt "
              f"+ @fit  (vyaz {row['vyaz']['line_count']} lines)")


if __name__ == "__main__":
    main()
