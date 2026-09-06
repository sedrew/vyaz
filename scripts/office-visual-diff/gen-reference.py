#!/usr/bin/env python3
"""
gen-reference.py — one slide per corpus case, a single auto-sized text box
holding the case's paragraphs, hairline-bordered so the export tells us exactly
where PowerPoint put the text block.

    python3 scripts/office-visual-diff/gen-reference.py

Reads packages/renderers/tests/cases/<group>/<case>/input.json directly (the
corpus is the single source of truth), keeps what a text box can represent, and
writes:

    _reference.pptx     open in PowerPoint → File ▸ Export ▸ PNG (one per slide)
    _manifest.json      slide → case name + box (left/top/width, pt) + skip list

Then: bun scripts/office-visual-diff/import-reference.ts <folder-of-pngs>

Units: every vyaz number is placed as a POINT here (Pt(fontSize), Pt(width), …).
`auto_size = SHAPE_TO_FIT_TEXT` + zero insets ⇒ the border rect in the PNG is
PowerPoint's own text-block box — measure-ref.ts reads it back.
"""
import json
import os
import sys

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
CORPUS = os.path.normpath(os.path.join(HERE, "../../packages/renderers/tests/cases"))

# Only these cases (VD_ONLY="a/b,c/d" overrides). The rich case covers "everything";
# writing-mode covers rotation + placement inside the frame.
ONLY = (os.environ.get("VD_ONLY") or
        "rich-text-v1/frame,"
        "writing-mode/rotate-180,writing-mode/rotate-270,"
        "writing-mode/sideways-lr,writing-mode/sideways-rl").split(",")

VERT = {"sideways-rl": "vert", "sideways-lr": "vert270"}  # CSS → a:bodyPr@vert

SLIDE_W_PT, SLIDE_H_PT = 960.0, 540.0
MARGIN_PT = 24.0
BORDER = RGBColor(0xFF, 0x00, 0x00)

ALIGN = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT,
         "center": PP_ALIGN.CENTER, "justify": PP_ALIGN.JUSTIFY}

DEFAULT_TEXT = {"fontFamily": "Arial", "fontSize": 12, "fontWeight": "normal",
                "fontStyle": "normal", "color": "#000000"}
DEFAULT_PARA = {"alignment": "left", "lineHeight": 1.15, "spaceBefore": 0, "spaceAfter": 0}


def hex_rgb(c, fallback=(0, 0, 0)):
    try:
        c = (c or "").lstrip("#")
        if len(c) == 3:
            c = "".join(ch * 2 for ch in c)
        return RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))
    except Exception:
        return RGBColor(*fallback)


def fam(f):
    return str(f[0]) if isinstance(f, list) and f else str(f or "Arial")


def skip_reason(inp):
    if inp.get("table"):
        return "TableFrame"
    fr = inp.get("frame")
    if not fr or not isinstance(fr.get("paragraphs"), list):
        return "no frame.paragraphs"
    if fr.get("columns"):
        return "multi-column"
    wm = fr.get("writingMode")
    if wm and wm not in ("horizontal-tb", "sideways-lr", "sideways-rl"):
        return f"writingMode {wm}"
    if (fr.get("autofit") or {}).get("enabled"):
        return "autofit"
    for p in fr["paragraphs"]:
        for r in p.get("children", []):
            if r.get("type") == "inline-box" or r.get("inlineWidget"):
                return "inline widget"
    return None  # list markers render as plain paragraphs (marker dropped)


def apply_transform(text, kind):
    if kind == "uppercase":
        return text.upper()
    if kind == "lowercase":
        return text.lower()
    if kind == "capitalize":  # vyaz: first letter of each word up, rest untouched
        import re
        return re.sub(r"(?<![^\W_])\w", lambda m: m.group().upper(), text)
    return text


def add_case(slide, frame):
    d = {**DEFAULT_TEXT, **(frame.get("defaultStyle") or {})}
    left, top = Pt(MARGIN_PT), Pt(MARGIN_PT)
    width = Pt(float(frame["width"]) if frame.get("width") else SLIDE_W_PT - 2 * MARGIN_PT)
    box = slide.shapes.add_textbox(left, top, width, Pt(40))
    tf = box.text_frame
    tf.word_wrap = bool(frame.get("wrap"))
    tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    paras = [p for p in frame["paragraphs"]
             if any(r.get("type") != "inline-box" and isinstance(r.get("text"), str) and r["text"]
                    for r in p.get("children", []))]
    for pi, p in enumerate(paras):
        ps = {**DEFAULT_PARA, **(p.get("style") or {})}
        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        para.alignment = ALIGN.get(ps.get("alignment", "left"), PP_ALIGN.LEFT)
        if ps.get("spaceBefore"):
            para.space_before = Pt(float(ps["spaceBefore"]))
        if ps.get("spaceAfter"):
            para.space_after = Pt(float(ps["spaceAfter"]))
        # Only override line spacing when the case explicitly set lineHeight,
        # and then as a MULTIPLE (spcPct) of PowerPoint's per-font "Single".
        # No lineHeight → leave it: PowerPoint uses the font's own line height.
        raw_lh = (p.get("style") or {}).get("lineHeight")
        if raw_lh is not None and abs(float(raw_lh) - 1.15) > 1e-6:  # 1.15 = vyaz default
            para.line_spacing = float(raw_lh)
        for r in p.get("children", []):
            if r.get("type") == "inline-box" or not isinstance(r.get("text"), str):
                continue
            s = {**d, **r}
            run = para.add_run()
            run.text = apply_transform(r["text"], s.get("textTransform"))
            f = run.font
            f.name = fam(s.get("fontFamily"))
            f.size = Pt(float(s.get("fontSize", 12)))
            fw = s.get("fontWeight")
            f.bold = fw == "bold" or (isinstance(fw, (int, float)) and fw >= 600)
            f.italic = s.get("fontStyle") == "italic"
            f.underline = bool(s.get("underline"))
            f.color.rgb = hex_rgb(s.get("color"))
            rPr = run._r.get_or_add_rPr()
            if s.get("strikethrough"):
                rPr.set("strike", "sngStrike")
            if s.get("letterSpacing"):
                rPr.set("spc", str(int(round(float(s["letterSpacing"]) * 100))))
            if s.get("script") == "super":
                rPr.set("baseline", "30000")
            elif s.get("script") == "sub":
                rPr.set("baseline", "-25000")

    if frame.get("rotation"):
        box.rotation = float(frame["rotation"])
    vert = VERT.get(frame.get("writingMode"))
    if vert:
        tf._txBody.bodyPr.set("vert", vert)

    box.fill.background()
    box.line.color.rgb = BORDER
    box.line.width = Pt(0.75)
    return [left / 12700, top / 12700, width / 12700]


def main():
    names = []
    for group in sorted(os.listdir(CORPUS)):
        gp = os.path.join(CORPUS, group)
        if group.startswith("_") or not os.path.isdir(gp):
            continue
        for c in sorted(os.listdir(gp)):
            if os.path.isfile(os.path.join(gp, c, "input.json")):
                names.append(f"{group}/{c}")

    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W_PT * 12700))
    prs.slide_height = Emu(int(SLIDE_H_PT * 12700))
    blank = prs.slide_layouts[6]

    slides, skipped = [], []
    for name in names:
        if name not in ONLY:
            continue
        inp = json.load(open(os.path.join(CORPUS, name, "input.json")))
        why = skip_reason(inp)
        if why:
            skipped.append({"name": name, "why": why})
            continue
        slide = prs.slides.add_slide(blank)
        box = add_case(slide, inp["frame"])
        slides.append({"slide": len(slides) + 1, "name": name, "box": box})

    prs.save(os.path.join(HERE, "_reference.pptx"))
    json.dump(
        {"slideWpt": SLIDE_W_PT, "slideHpt": SLIDE_H_PT, "marginPt": MARGIN_PT,
         "borderHex": "FF0000", "slides": slides, "skipped": skipped},
        open(os.path.join(HERE, "_manifest.json"), "w"), indent=2,
    )
    print(f"wrote _reference.pptx ({len(slides)} slides) + _manifest.json  (skipped {len(skipped)})")
    by = {}
    for s in skipped:
        by.setdefault(s["why"], []).append(s["name"])
    for why, ns in sorted(by.items(), key=lambda kv: -len(kv[1])):
        print(f"  {why:<24} {len(ns):>2}  ({', '.join(ns[:3])}{', …' if len(ns) > 3 else ''})")
    print("\n→ open _reference.pptx in PowerPoint, File ▸ Export ▸ PNG (one per slide),")
    print("  then: bun scripts/office-visual-diff/import-reference.ts <png-folder>")


if __name__ == "__main__":
    sys.exit(main())
